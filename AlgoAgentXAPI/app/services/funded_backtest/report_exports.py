from __future__ import annotations

import io
import json
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from ...utils.timezone import format_kolkata_datetime


def _num(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except Exception:
        return None


def _money(value: Any, currency: str = "USD") -> str:
    number = _num(value)
    if number is None:
        return "—"
    symbol = {"USD": "$", "EUR": "€", "GBP": "£", "INR": "₹"}.get(str(currency).upper(), f"{currency} ")
    return f"{symbol}{number:,.2f}"


def _pct(value: Any) -> str:
    number = _num(value)
    return "—" if number is None else f"{number * 100:.2f}%"


def _ist(value: Any) -> str:
    return format_kolkata_datetime(value, fallback="—", include_timezone=True)


def build_funded_excel(report: dict[str, Any], phases: list[dict], trades: list[dict], days: list[dict], events: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    currency = ((report.get("account_snapshot") or {}).get("profile") or {}).get("account_currency") or "USD"
    summary = report.get("summary") or {}
    payout = summary.get("payout") or {}
    rows = [
        ("Funded Backtest Report", ""),
        ("Status", report.get("status")),
        ("Account", ((report.get("account_snapshot") or {}).get("profile") or {}).get("name")),
        ("Strategy", summary.get("source_strategy_name")),
        ("Market", f"{summary.get('source_instrument_symbol') or report.get('instrument_id')} · {report.get('timeframe')}"),
        ("Initial Capital", _money(report.get("initial_capital"), currency)),
        ("Final Balance", _money(report.get("final_balance"), currency)),
        ("Return", _pct(summary.get("return_pct"))),
        ("Max DD", _pct(summary.get("max_drawdown_pct"))),
        ("Trading Days", report.get("trading_days")),
        ("Qualifying Days", report.get("qualifying_days")),
        ("Calendar Days", report.get("calendar_days")),
        ("Estimated Payout", _money(payout.get("estimated_payout_amount"), currency)),
        ("Consistency", _pct(payout.get("consistency_pct"))),
        ("Engine Version", summary.get("engine_version") or report.get("rule_engine_version")),
        ("Drawdown Evaluation", summary.get("drawdown_evaluation_mode")),
    ]
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=16)
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 48

    def add_sheet(name: str, columns: list[tuple[str, str]], records: list[dict]):
        sh = wb.create_sheet(name)
        sh.append([label for label, _ in columns])
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="4C1D95")
            c.alignment = Alignment(horizontal="center")
        datetime_keys = {"starting_at", "ending_at", "passed_at", "failed_at", "entry_time", "exit_time", "event_timestamp", "created_at", "updated_at"}
        for record in records:
            sh.append([_ist(record.get(key)) if key in datetime_keys and record.get(key) else record.get(key) for _, key in columns])
        sh.freeze_panes = "A2"
        for idx, (label, _) in enumerate(columns, 1):
            sh.column_dimensions[get_column_letter(idx)].width = min(max(len(label) + 4, 14), 32)
        return sh

    add_sheet("Phases", [
        ("Phase", "phase_number"), ("Name", "phase_name"), ("Status", "phase_status"),
        ("Start", "starting_at"), ("End", "ending_at"), ("Start Balance", "starting_balance"),
        ("End Balance", "ending_balance"), ("Target %", "target_pct"), ("Target Amount", "target_amount"),
        ("Trading Days", "trading_days"), ("Qualifying Days", "qualifying_days"),
        ("Max DD", "maximum_drawdown"), ("Worst Daily DD", "worst_daily_drawdown"),
        ("Passed At", "passed_at"), ("Failed At", "failed_at"), ("Failure Reason", "failure_reason"),
    ], phases)
    add_sheet("Trades", [
        ("#", "trade_number"), ("Phase ID", "funded_phase_id"), ("Entry", "entry_time"), ("Exit", "exit_time"),
        ("Side", "side"), ("Risk Tier", "selected_risk_tier_name"), ("Requested Risk %", "requested_risk_pct"),
        ("Effective Risk %", "effective_risk_pct"), ("Actual Risk", "actual_risk_amount"), ("Lot", "calculated_lot_size"),
        ("Quantity", "calculated_quantity"), ("Entry Price", "entry_price"), ("SL", "stop_loss"), ("TP", "target"),
        ("Exit Price", "exit_price"), ("PnL", "pnl"), ("R", "r_multiple"), ("Balance Before", "balance_before_trade"),
        ("Balance After", "balance_after_trade"), ("Daily DD %", "daily_dd_pct_used"), ("Max DD %", "max_dd_pct_used"),
        ("State", "account_state_after_trade"),
    ], trades)
    add_sheet("Daily", [
        ("Date", "trading_date"), ("Start Balance", "start_balance"), ("End Balance", "end_balance"),
        ("PnL", "day_pnl"), ("Return %", "day_return_pct"), ("Trades", "trades_count"), ("Wins", "winning_trades"),
        ("Losses", "losing_trades"), ("Daily DD %", "daily_drawdown_pct"), ("Qualifying", "qualifying_day"),
        ("Threshold", "qualifying_profit_threshold"), ("Consistency %", "consistency_pct"), ("State", "state"),
    ], days)
    risk_records = []
    for name, count in (summary.get("risk_tier_distribution") or {}).items():
        risk_records.append({"tier": name, "trades": count})
    add_sheet("Risk Analysis", [("Tier", "tier"), ("Trades", "trades")], risk_records)
    add_sheet("Events", [
        ("Timestamp", "event_timestamp"), ("Type", "event_type"), ("Title", "event_title"), ("Message", "message"), ("Metadata", "metadata_json")
    ], [{**x, "metadata_json": json.dumps(x.get("metadata_json") or {}, default=str)} for x in events])
    rules = wb.create_sheet("Rules Snapshot")
    rules["A1"] = "Account Snapshot"
    rules["A2"] = json.dumps(report.get("account_snapshot") or {}, indent=2, default=str)
    rules["A4"] = "Risk Plan Snapshot"
    rules["A5"] = json.dumps(report.get("risk_plan_snapshot") or {}, indent=2, default=str)
    rules["A7"] = "Runtime Snapshot"
    rules["A8"] = json.dumps(report.get("runtime_snapshot") or {}, indent=2, default=str)
    rules.column_dimensions["A"].width = 110
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def build_funded_pdf(report: dict[str, Any], phases: list[dict], days: list[dict], events: list[dict]) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=16*mm, bottomMargin=16*mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("FundedTitle", parent=styles["Title"], textColor=colors.HexColor("#4C1D95"), spaceAfter=8)
    h2 = ParagraphStyle("FundedH2", parent=styles["Heading2"], textColor=colors.HexColor("#4C1D95"), spaceBefore=10, spaceAfter=6)
    story = [Paragraph("AlgoAgentX Funded Backtest Report", title)]
    summary = report.get("summary") or {}
    profile = ((report.get("account_snapshot") or {}).get("profile") or {})
    payout = summary.get("payout") or {}
    currency = profile.get("account_currency") or "USD"
    story.append(Paragraph(f"<b>{report.get('status') or '—'}</b> · {profile.get('name') or 'Funded Account'} · {summary.get('source_strategy_name') or report.get('strategy_id')}", styles["BodyText"]))
    story.append(Spacer(1, 8))
    data = [
        ["Initial", _money(report.get("initial_capital"), currency), "Final", _money(report.get("final_balance"), currency)],
        ["Return", _pct(summary.get("return_pct")), "Max DD", _pct(summary.get("max_drawdown_pct"))],
        ["Trading Days", str(report.get("trading_days") or 0), "Qualifying Days", str(report.get("qualifying_days") or 0)],
        ["Calendar Days", str(report.get("calendar_days") or 0), "Estimated Payout", _money(payout.get("estimated_payout_amount"), currency)],
    ]
    table = Table(data, colWidths=[34*mm, 45*mm, 34*mm, 45*mm])
    table.setStyle(TableStyle([("GRID", (0,0), (-1,-1), .4, colors.HexColor("#D1D5DB")), ("BACKGROUND", (0,0), (0,-1), colors.HexColor("#F3E8FF")), ("BACKGROUND", (2,0), (2,-1), colors.HexColor("#F3E8FF")), ("FONTNAME", (0,0), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), 9), ("VALIGN", (0,0), (-1,-1), "MIDDLE")]))
    story.append(table)
    story.append(Paragraph("Account & Rules", h2))
    story.append(Paragraph(f"Challenge type: {profile.get('challenge_type') or '—'} · Account size: {_money(profile.get('account_size'), currency)} · Risk mode: {summary.get('risk_mode') or '—'} · Drawdown evaluation: {summary.get('drawdown_evaluation_mode') or '—'} · Engine: {summary.get('engine_version') or report.get('rule_engine_version') or '—'}", styles["BodyText"]))
    if phases:
        story.append(Paragraph("Phases", h2))
        pdata = [["Phase", "Status", "Start", "End", "Target", "Trading", "Qualifying"]]
        for p in phases:
            pdata.append([str(p.get("phase_number") or "—"), str(p.get("phase_status") or "—"), _money(p.get("starting_balance"), currency), _money(p.get("ending_balance"), currency), _pct(p.get("target_pct")), str(p.get("trading_days") or 0), str(p.get("qualifying_days") or 0)])
        pt = Table(pdata, repeatRows=1)
        pt.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#4C1D95")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .3, colors.HexColor("#D1D5DB")), ("FONTSIZE", (0,0), (-1,-1), 8)]))
        story.append(pt)
    if payout:
        story.append(Paragraph("Payout", h2))
        story.append(Paragraph(f"Gross eligible profit: {_money(payout.get('gross_eligible_profit'), currency)} · Profit split: {_pct(payout.get('profit_split_pct'))} · Estimated payout: {_money(payout.get('estimated_payout_amount'), currency)} · Consistency: {_pct(payout.get('consistency_pct'))}", styles["BodyText"]))
    if days:
        best = max(days, key=lambda x: _num(x.get("day_pnl")) or float("-inf"))
        worst = min(days, key=lambda x: _num(x.get("day_pnl")) or float("inf"))
        story.append(Paragraph("Daily Performance", h2))
        story.append(Paragraph(f"Best day: {best.get('trading_date')} {_money(best.get('day_pnl'), currency)} · Worst day: {worst.get('trading_date')} {_money(worst.get('day_pnl'), currency)}", styles["BodyText"]))
    if events:
        story.append(Paragraph("Key Rule Events", h2))
        for event in events[:20]:
            story.append(Paragraph(f"<b>{event.get('event_type')}</b> — {event.get('event_title')} ({_ist(event.get('event_timestamp'))})", styles["BodyText"]))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Full trade-by-trade audit is available in the funded Excel and CSV exports.", styles["Italic"]))
    doc.build(story)
    return output.getvalue()
